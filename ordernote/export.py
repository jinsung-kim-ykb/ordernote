"""출력: 배송표 엑셀, 거래처 확인서, 확인큐 CSV."""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import Item, ResolvedLine

DISTRIBUTOR = "새벽들유통"
YELLOW = PatternFill("solid", fgColor="FFF59D")
BOLD = Font(bold=True)


def raw_text(r: ResolvedLine) -> str:
    l = r.line
    qty = f" {l.수량:g}{l.단위 or ''}" if l.수량 is not None else ""
    return f"{l.품목명_원문}{qty}"


def fmt_item(item: Item, qty: float) -> str:
    if item.단위 == "kg":
        return f"{item.품목명} {qty:g}kg"
    if item.단위 == "단":
        return f"{item.품목명} {item.규격}×{qty:g}"
    return f"{item.품목명} {qty:g}{item.단위}"


def _sheet_title(name: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "", name)[:31]


def _write_rows(ws, header: list[str], rows: list[list], highlight: list[bool] | None = None):
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD
    for i, row in enumerate(rows):
        ws.append(row)
        if highlight and highlight[i]:
            for c in ws[ws.max_row]:
                c.fill = YELLOW
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(8, width * 1.6 + 2), 60)
    ws.freeze_panes = "A2"


def write_xlsx(resolved: list[ResolvedLine], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "피킹리스트"
    totals: dict[str, float] = defaultdict(float)
    per_client: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    items: dict[str, Item] = {}
    for r in resolved:
        if r.status == "정상":
            code = r.item.품목코드
            items[code] = r.item
            totals[code] += r.수량
            per_client[code][r.order.거래처] += r.수량
    rows = [
        [c, items[c].품목명, items[c].규격, items[c].단위, totals[c], len(per_client[c]),
         ", ".join(f"{k} {v:g}" for k, v in per_client[c].items())]
        for c in sorted(totals)
    ]
    _write_rows(ws, ["품목코드", "품목명", "규격", "단위", "합계수량", "거래처수", "거래처별"], rows)
    pending = sum(r.status == "확인필요" for r in resolved)
    if pending:
        ws.append([])
        ws.append([f"※ 확인필요 {pending}줄은 피킹 제외 — 확인큐.csv 참고"])
        ws[ws.max_row][0].fill = YELLOW

    by_client: dict[str, list[ResolvedLine]] = defaultdict(list)
    for r in resolved:
        by_client[r.order.거래처].append(r)
    for client, rs in by_client.items():
        ws = wb.create_sheet(_sheet_title(client))
        rows = [
            [r.item.품목코드 if r.item else "", r.item.품목명 if r.item else "", r.item.규격 if r.item else "",
             r.수량, r.단위, r.status, r.사유, raw_text(r), r.line.메모 or "",
             f"#{r.order.순번} {r.order.수신시각:%H:%M}"]
            for r in rs
        ]
        _write_rows(ws, ["품목코드", "품목명", "규격", "수량", "단위", "상태", "사유", "원문", "메모", "주문"],
                    rows, [r.status == "확인필요" for r in rs])
    wb.save(path)


def write_confirmations(resolved: list[ResolvedLine], out_dir: Path) -> int:
    out_dir.mkdir(parents=True, exist_ok=True)
    by_client: dict[str, list[ResolvedLine]] = defaultdict(list)
    for r in resolved:
        by_client[r.order.거래처].append(r)
    for client, rs in by_client.items():
        ok: dict[str, float] = defaultdict(float)
        items: dict[str, Item] = {}
        dates = set()
        for r in rs:
            if r.status == "정상":
                ok[r.item.품목코드] += r.수량
                items[r.item.품목코드] = r.item
                if r.line.납품일:
                    dates.add(r.line.납품일)
        pending = [raw_text(r) for r in rs if r.status == "확인필요"]
        when = "·".join(f"{date.fromisoformat(d):%m/%d}".lstrip("0").replace("/0", "/") for d in sorted(dates)) or "다음"
        text = [f"[{DISTRIBUTOR}] 주문 확인"]
        if ok:
            text.append(f"{client}님, {when} 납품 주문 확인드립니다. " + ", ".join(fmt_item(items[c], q) for c, q in ok.items()))
        else:
            text.append(f"{client}님, 주문 잘 받았습니다.")
        if pending:
            text.append(f"※ 확인 중: {', '.join(pending)} — 품목/수량 확인 후 다시 연락드리겠습니다.")
        text.append("변경은 이 채널로 회신 주세요.")
        (out_dir / f"{client}.txt").write_text("\n".join(text) + "\n", encoding="utf-8")
    return len(by_client)


def write_queue(resolved: list[ResolvedLine], path: Path) -> int:
    rows = [r for r in resolved if r.status == "확인필요"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["거래처", "수신시각", "원문", "후보1", "후보2", "후보3", "사유", "메시지 전문"])
        for r in rows:
            c = (r.후보 + ["", "", ""])[:3]
            w.writerow([r.order.거래처, f"{r.order.수신시각:%Y-%m-%d %H:%M}", raw_text(r), *c, r.사유,
                        r.order.본문.replace("\n", " / ")])
    return len(rows)
